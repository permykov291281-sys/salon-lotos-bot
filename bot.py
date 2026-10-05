"""Настена — ИИ-администратор салона красоты «ЛОТОС» для сообщества ВКонтакте.

VK Long Poll -> нейросеть (Proxy API, function calling) -> запись в Excel data/zapisi_LOTOS.xlsx.
Свободное время считается по графику мастеров (лист «Мастера») и уже сделанным записям (лист «Записи»).
Запуск в VK: python3 bot.py        Проверка в консоли: python3 bot.py --cli
"""
import json, os, random, re, sys, time, uuid, urllib.error, urllib.parse, urllib.request
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE / "vendor"))  # необязательно: openpyxl можно положить в vendor/ или поставить через pip
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
import kb_data as KB  # noqa: E402

for line in (HERE / ".env").read_text().splitlines() if (HERE / ".env").exists() else []:
    if "=" in line and not line.startswith("#"):
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip())

VK_TOKEN = os.environ.get("VK_TOKEN", "")
PROXYAPI_KEY = os.environ["PROXYAPI_KEY"]
GROUP_ID = int(os.environ.get("VK_GROUP_ID", "0") or 0)
MODEL = os.environ.get("MODEL", "gpt-4.1")
AI_URL = os.environ.get("AI_URL", "https://api.proxyapi.ru/openai/v1/chat/completions")  # любой OpenAI-совместимый endpoint
ADMIN_IDS = [int(x) for x in os.environ.get("ADMIN_VK_IDS", "").split(",") if x.strip()]
TZ = ZoneInfo("Asia/Yekaterinburg")
HISTORY_LIMIT = 30
SLOT_STEP = 30          # шаг сетки времени, минут
MIN_LEAD_MIN = 60       # не записывать раньше чем через час от «сейчас»
MAX_DAYS = 14           # насколько вперёд можно смотреть расписание

DATA = HERE / "data"; DATA.mkdir(exist_ok=True)
XLSX = DATA / "zapisi_LOTOS.xlsx"
HISTORY_FILE = DATA / "history.json"
history = json.loads(HISTORY_FILE.read_text()) if HISTORY_FILE.exists() else {}
DAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]
BOOK_COLS = ["ID", "Создано", "Клиент", "Телефон", "VK", "Услуга", "Мастер", "Дата", "Время",
             "Длительность, мин", "Цена, ₽", "Статус", "Комментарий"]
current_uid = 0
confirmed_now = False
CONFIRM = re.compile(r"^(да|ага|угу|ок|окей|хорошо|давай|давайте|конечно|верно|всё верно|все верно|подтверждаю|записывай|записывайте|\+)\b|^да[,! .]|записыв", re.I)


def log(*a):
    print(datetime.now(TZ).strftime("%d.%m %H:%M:%S"), *a, flush=True)


# ---------- база знаний ----------
def knowledge_base():
    L = ["О салоне:"] + [f"- {k}: {v}" for k, v in KB.SALON]
    L += ["", "Услуги (название — длительность — цена — описание):"]
    for cat in dict.fromkeys(s[0] for s in KB.SERVICES):
        L.append(f"{cat}:")
        L += [f"- {n} — {d} мин — {p} ₽ — {desc}" for c, n, d, p, desc in KB.SERVICES if c == cat]
    L += ["", "Мастера (имя — направление):"] + [f"- {m} — {s}" for m, s, *_ in KB.MASTERS]
    L += ["", "Частые вопросы:"] + [f"- {q} {a}" for q, a in KB.FAQ]
    L += ["", "Стандарты салона:"] + [f"- {s}" for s in KB.STANDARDS]
    L += ["", "Тон:"] + [f"- {s}" for s in KB.TONE]
    return "\n".join(L)


def system_prompt():
    now = datetime.now(TZ)
    return ((HERE / "system_prompt.md").read_text()
            .replace("{TODAY}", f"{now:%Y-%m-%d}, {DAYS[now.weekday()]}")
            .replace("{NOW}", f"{now:%H:%M}")
            .replace("{KNOWLEDGE_BASE}", knowledge_base()))


# ---------- Excel ----------
def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="B5577A")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def ensure_workbook():
    if XLSX.exists():
        return
    wb = Workbook(); ws = wb.active; ws.title = "Записи"; ws.append(BOOK_COLS)
    for col, w in zip("ABCDEFGHIJKLM", [10, 17, 16, 16, 28, 30, 18, 12, 8, 10, 10, 12, 30]):
        ws.column_dimensions[col].width = w
    style_header(ws)
    ws = wb.create_sheet("Лиды"); ws.append(["Создано", "Имя", "Телефон", "VK", "Интерес / комментарий"]); style_header(ws)
    for col, w in zip("ABCDE", [17, 16, 16, 28, 50]): ws.column_dimensions[col].width = w
    ws = wb.create_sheet("Мастера"); ws.append(["Мастер", "Направления", "Рабочие дни", "Начало", "Конец"]); style_header(ws)
    for m in KB.MASTERS: ws.append(list(m))
    for col, w in zip("ABCDE", [20, 18, 22, 10, 10]): ws.column_dimensions[col].width = w
    ws = wb.create_sheet("Услуги"); ws.append(["Категория", "Услуга", "Длительность, мин", "Цена, ₽", "Описание"]); style_header(ws)
    for s in KB.SERVICES: ws.append(list(s))
    for col, w in zip("ABCDE", [12, 34, 12, 10, 60]): ws.column_dimensions[col].width = w
    wb.save(XLSX)


def rows(sheet):
    ensure_workbook()
    ws = load_workbook(XLSX, data_only=True)[sheet]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]


def append_row(sheet, values):
    ensure_workbook()
    wb = load_workbook(XLSX); wb[sheet].append(values); wb.save(XLSX)


def set_status(booking_id, status):
    wb = load_workbook(XLSX); ws = wb["Записи"]
    for r in ws.iter_rows(min_row=2):
        if str(r[0].value) == booking_id:
            r[11].value = status; wb.save(XLSX); return True
    return False


# ---------- расписание ----------
def find_service(name):
    name = (name or "").lower().strip()
    exact = [s for s in KB.SERVICES if s[1].lower() == name]
    if exact:
        return exact[0]
    part = [s for s in KB.SERVICES if name and (name in s[1].lower() or s[1].lower() in name)]
    return part[0] if len(part) == 1 else None


def masters_for(category, master_hint=None):
    res = []
    for m in rows("Мастера"):
        if category not in str(m["Направления"]):
            continue
        if master_hint and master_hint.lower().split()[0] not in str(m["Мастер"]).lower():
            continue
        res.append(m)
    return res


def hm(s):
    s = str(s)[:5]; h, m = s.split(":"); return int(h) * 60 + int(m)


def busy(master, date):
    out = []
    for b in rows("Записи"):
        if b["Мастер"] == master and str(b["Дата"]) == date and b["Статус"] == "записан":
            st = hm(b["Время"]); out.append((st, st + int(b["Длительность, мин"])))
    return out


def free_times(m, date_str, duration):
    d = datetime.strptime(date_str, "%Y-%m-%d").date()
    if DAYS[d.weekday()] not in str(m["Рабочие дни"]).replace(" ", "").split(","):
        return []
    start, end = hm(m["Начало"]), hm(m["Конец"])
    now = datetime.now(TZ)
    earliest = (now.hour * 60 + now.minute + MIN_LEAD_MIN) if d == now.date() else 0
    taken = busy(m["Мастер"], date_str)
    res = []
    t = start
    while t + duration <= end:
        if t >= earliest and all(t + duration <= a or t >= b for a, b in taken):
            res.append(f"{t // 60:02d}:{t % 60:02d}")
        t += SLOT_STEP
    return res


# ---------- инструменты для нейросети ----------
def tool_find_free_slots(service, date_from, date_to=None, master=None):
    s = find_service(service)
    if not s:
        return {"error": "Услуга не найдена. Есть: " + ", ".join(x[1] for x in KB.SERVICES)}
    today = datetime.now(TZ).date()
    try:
        d1 = max(datetime.strptime(date_from, "%Y-%m-%d").date(), today)
        d2 = datetime.strptime(date_to or date_from, "%Y-%m-%d").date()
    except ValueError:
        return {"error": "Даты нужны в формате ГГГГ-ММ-ДД"}
    d2 = min(max(d2, d1), d1 + timedelta(days=6), today + timedelta(days=MAX_DAYS))
    ms = masters_for(s[0], master)
    if not ms:
        return {"error": f"Нет мастера «{master}» по направлению {s[0]}. Мастера: " + ", ".join(m["Мастер"] for m in masters_for(s[0]))}
    out = []
    d = d1
    while d <= d2:
        for m in ms:
            times = free_times(m, d.isoformat(), s[2])
            if times:
                out.append({"date": d.isoformat(), "weekday": DAYS[d.weekday()], "master": m["Мастер"], "free": times})
        d += timedelta(days=1)
    return {"service": s[1], "duration_min": s[2], "price": s[3], "slots": out or "свободного времени в эти даты нет"}


def tool_create_booking(client_name, phone, service, master, date, time, comment=""):
    s = find_service(service)
    if not s:
        return {"error": "Услуга не найдена"}
    ms = masters_for(s[0], master)
    if len(ms) != 1:
        return {"error": "Уточните мастера: " + ", ".join(m["Мастер"] for m in masters_for(s[0]))}
    digits = re.sub(r"\D", "", phone or "")
    if len(digits) < 10:
        return {"error": "Нужен телефон клиента (10–11 цифр)"}
    m = ms[0]
    same = [b for b in rows("Записи") if b["Статус"] == "записан" and b["Мастер"] == m["Мастер"] and str(b["Дата"]) == date
            and str(b["Время"])[:5] == time and re.sub(r"\D", "", str(b["Телефон"]))[-10:] == digits[-10:]]
    if same:
        return {"ok": True, "already_booked": True, "booking_id": same[0]["ID"], "service": s[1], "master": m["Мастер"], "date": date, "time": time}
    if not confirmed_now:
        return {"error": "Клиент ещё не подтвердил. Сначала коротко повтори запись (услуга, мастер, дата, время, цена) и спроси «Записываю?», дождись ответа «да»."}
    try:
        if time not in free_times(m, date, s[2]):
            return {"error": f"Время {date} {time} у мастера {m['Мастер']} уже занято или нерабочее. Вызови find_free_slots и предложи другое."}
    except ValueError:
        return {"error": "Дата ГГГГ-ММ-ДД, время ЧЧ:ММ"}
    bid = uuid.uuid4().hex[:6].upper()
    append_row("Записи", [bid, datetime.now(TZ).strftime("%Y-%m-%d %H:%M"), client_name, phone,
                          f"https://vk.com/id{current_uid}" if current_uid else "", s[1], m["Мастер"], date, time,
                          s[2], s[3], "записан", comment])
    notify_admins(f"🌸 Новая запись {bid}\n{client_name}, {phone}\n{s[1]} — {m['Мастер']}\n"
                  f"{date} ({DAYS[datetime.strptime(date, '%Y-%m-%d').weekday()]}) в {time}, {s[2]} мин, {s[3]} ₽"
                  + (f"\nКомментарий: {comment}" if comment else "")
                  + (f"\nVK: https://vk.com/id{current_uid}" if current_uid else ""))
    log("запись", bid, s[1], m["Мастер"], date, time)
    return {"ok": True, "booking_id": bid, "service": s[1], "master": m["Мастер"], "date": date, "time": time, "price": s[3]}


def tool_find_my_bookings(phone):
    digits = re.sub(r"\D", "", phone or "")[-10:]
    found = [{"booking_id": b["ID"], "service": b["Услуга"], "master": b["Мастер"], "date": str(b["Дата"]), "time": str(b["Время"])}
             for b in rows("Записи") if b["Статус"] == "записан" and re.sub(r"\D", "", str(b["Телефон"]))[-10:] == digits]
    return {"bookings": found or "записей не найдено"}


def tool_cancel_booking(phone, date=None):
    digits = re.sub(r"\D", "", phone or "")[-10:]
    mine = [b for b in rows("Записи") if b["Статус"] == "записан" and re.sub(r"\D", "", str(b["Телефон"]))[-10:] == digits
            and (not date or str(b["Дата"]) == date)
            and (not b["VK"] or not current_uid or b["VK"].endswith(f"id{current_uid}"))]  # чужие записи не отменяем
    if not mine:
        return {"error": "Активных записей с этим телефоном не найдено"}
    if len(mine) > 1:
        return {"error": "Несколько записей, уточни дату", "bookings": [f"{b['Дата']} {b['Время']} {b['Услуга']}" for b in mine]}
    b = mine[0]
    set_status(b["ID"], "отменён")
    notify_admins(f"❌ Отмена записи {b['ID']}: {b['Клиент']}, {b['Услуга']} — {b['Мастер']}, {b['Дата']} {b['Время']}")
    return {"ok": True, "cancelled": f"{b['Дата']} {b['Время']} {b['Услуга']} — {b['Мастер']}"}


def tool_save_lead(client_name, phone, interest):
    append_row("Лиды", [datetime.now(TZ).strftime("%Y-%m-%d %H:%M"), client_name, phone,
                        f"https://vk.com/id{current_uid}" if current_uid else "", interest])
    notify_admins(f"📝 Контакт без записи: {client_name}, {phone}\n{interest}")
    return {"ok": True}


TOOLS = {"find_free_slots": tool_find_free_slots, "create_booking": tool_create_booking,
         "find_my_bookings": tool_find_my_bookings, "cancel_booking": tool_cancel_booking, "save_lead": tool_save_lead}


def fn(name, desc, props, req):
    return {"type": "function", "function": {"name": name, "description": desc,
            "parameters": {"type": "object", "properties": {k: {"type": "string", "description": v} for k, v in props.items()}, "required": req}}}


TOOL_SCHEMAS = [
    fn("find_free_slots", "Свободное время мастеров для услуги в диапазоне дат (не больше 7 дней).",
       {"service": "Точное название услуги из базы", "date_from": "ГГГГ-ММ-ДД", "date_to": "ГГГГ-ММ-ДД", "master": "Имя мастера, если клиент хочет конкретного"},
       ["service", "date_from"]),
    fn("create_booking", "Записать клиента. Только после явного согласия, с именем и телефоном, на время из find_free_slots.",
       {"client_name": "Имя", "phone": "Телефон", "service": "Название услуги", "master": "Имя мастера",
        "date": "ГГГГ-ММ-ДД", "time": "ЧЧ:ММ", "comment": "Пожелания клиента"},
       ["client_name", "phone", "service", "master", "date", "time"]),
    fn("find_my_bookings", "Найти активные записи клиента по телефону.", {"phone": "Телефон"}, ["phone"]),
    fn("cancel_booking", "Отменить запись клиента по телефону (и дате, если записей несколько).",
       {"phone": "Телефон", "date": "ГГГГ-ММ-ДД"}, ["phone"]),
    fn("save_lead", "Сохранить контакт клиента, который пока не записался (или жалобу).",
       {"client_name": "Имя", "phone": "Телефон", "interest": "Что интересует / комментарий"}, ["client_name", "phone", "interest"]),
]


# ---------- сеть ----------
def http_post(url, data=None, headers=None, body=None, timeout=60, tries=4):
    if data is not None:
        body = urllib.parse.urlencode(data).encode()
    for attempt in range(tries):
        try:
            req = urllib.request.Request(url, data=body, headers=headers or {})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return json.loads(r.read().decode())
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            if attempt == tries - 1:
                raise RuntimeError(f"{urllib.parse.urlsplit(url).netloc}: {e}")
            time.sleep(1 + attempt * 2)


def ask_ai(messages):
    return http_post(AI_URL,
                     headers={"Authorization": f"Bearer {PROXYAPI_KEY}", "Content-Type": "application/json"},
                     body=json.dumps({"model": MODEL, "temperature": 0.5, "messages": messages,
                                      "tools": TOOL_SCHEMAS}).encode(), timeout=90)["choices"][0]["message"]


def vk(method, **params):
    params.update(access_token=VK_TOKEN, v="5.199")
    res = http_post(f"https://api.vk.com/method/{method}", params, timeout=20)
    if "error" in res:
        raise RuntimeError(f"VK {method}: {res['error'].get('error_msg')}")
    return res["response"]


def send(peer_id, text, attachment=None):
    if CLI:
        print(("\n[АДМИНУ] " if peer_id in ADMIN_IDS and peer_id != current_uid else "\nНастена: ") + text); return
    p = dict(peer_id=peer_id, message=text, random_id=random.randint(1, 2**31))
    if attachment: p["attachment"] = attachment
    vk("messages.send", **p)


def notify_admins(text):
    if CLI:
        print("\n[УВЕДОМЛЕНИЕ АДМИНУ] " + text); return
    for a in ADMIN_IDS:
        try: send(a, text)
        except Exception as e: log("не удалось уведомить", a, e)


def send_table(peer_id):
    """Отправить администратору файл Excel с записями."""
    ensure_workbook()
    up = vk("docs.getMessagesUploadServer", peer_id=peer_id, type="doc")["upload_url"]
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"zapisi_LOTOS.xlsx\"\r\n"
            f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() \
        + XLSX.read_bytes() + f"\r\n--{boundary}--\r\n".encode()
    f = http_post(up, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}, body=body)["file"]
    doc = vk("docs.save", file=f, title=f"Записи ЛОТОС {datetime.now(TZ):%d.%m %H-%M}")["doc"]
    send(peer_id, "Таблица записей 👇", attachment=f"doc{doc['owner_id']}_{doc['id']}")


def clean(text):
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text or "")
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*[\*•]\s+", "— ", text, flags=re.M)
    return text.replace("`", "").strip()


# ---------- диалог ----------
def handle(uid, text):
    global current_uid
    current_uid = uid
    key = str(uid); low = text.strip().lower()
    if low in ("/reset", "/сброс"):
        history.pop(key, None); send(uid, "История диалога очищена."); return
    if low.lstrip("/") in ("id", "ид", "айди"):
        send(uid, f"Ваш VK ID: {uid}"); return
    if uid in ADMIN_IDS and low in ("/записи", "/таблица"):
        send_table(uid); return
    msgs = history.setdefault(key, [])
    msgs.append({"role": "user", "content": text})
    if not CLI:
        try: vk("messages.setActivity", peer_id=uid, type="typing")
        except Exception: pass
    global confirmed_now
    prev_bot = next((m["content"] for m in reversed(msgs[:-1]) if m["role"] == "assistant"), "")
    confirmed_now = "?" in prev_bot and bool(CONFIRM.search(low))
    convo = [{"role": "system", "content": system_prompt()}] + msgs[-HISTORY_LIMIT:]
    answer = None
    for _ in range(6):  # цикл вызова инструментов
        msg = ask_ai(convo)
        calls = msg.get("tool_calls") or []
        if not calls:
            answer = msg.get("content"); break
        convo.append({"role": "assistant", "content": msg.get("content"), "tool_calls": calls})
        for c in calls:
            try:
                args = json.loads(c["function"]["arguments"] or "{}")
                result = TOOLS[c["function"]["name"]](**args)
            except Exception as e:
                result = {"error": str(e)}
            log("tool", c["function"]["name"], json.dumps(args, ensure_ascii=False)[:150], "->", json.dumps(result, ensure_ascii=False)[:150])
            convo.append({"role": "tool", "tool_call_id": c["id"], "content": json.dumps(result, ensure_ascii=False)})
    answer = clean(answer) or "Секундочку, уточню и вернусь 🙏"
    msgs.append({"role": "assistant", "content": answer})
    history[key] = msgs[-HISTORY_LIMIT:]
    HISTORY_FILE.write_text(json.dumps(history, ensure_ascii=False))
    send(uid, answer)


def safe_handle(uid, text):
    try:
        handle(uid, text)
    except Exception as e:
        log("ошибка обработки", uid, ":", e)
        try: send(uid, "Ой, у меня что-то подвисло 🙈 Напишите, пожалуйста, ещё раз через минутку.")
        except Exception: pass


def catch_up():
    try:
        convs = vk("messages.getConversations", group_id=GROUP_ID, filter="unanswered", count=20)["items"]
    except Exception as e:
        log("catch_up:", e); return
    for c in convs:
        uid = c["conversation"]["peer"]["id"]
        if uid <= 0 or time.time() - c["last_message"]["date"] > 3600:
            continue
        items = vk("messages.getHistory", peer_id=uid, group_id=GROUP_ID, count=20)["items"]
        unanswered = []
        for m in items:
            if m["out"]: break
            if m.get("text"): unanswered.insert(0, m["text"])
        if unanswered:
            log("догоняю", uid); safe_handle(uid, "\n".join(unanswered))


def main():
    ensure_workbook()
    log(f"Настена запущена: сообщество {GROUP_ID}, модель {MODEL}, админов {len(ADMIN_IDS)}")
    catch_up()
    srv, ts = None, None
    while True:
        try:
            if srv is None:
                srv = vk("groups.getLongPollServer", group_id=GROUP_ID); ts = ts or srv["ts"]
            res = http_post(srv["server"], {"act": "a_check", "key": srv["key"], "ts": ts, "wait": 25}, timeout=35)
            if "failed" in res:
                if res["failed"] == 1: ts = res["ts"]
                else:
                    srv = None
                    if res["failed"] == 3: ts = None
                continue
            ts = res["ts"]
            for ev in res.get("updates", []):
                if ev["type"] == "message_new":
                    m = ev["object"]["message"]
                    if m.get("from_id", 0) > 0 and m.get("text"):
                        log("сообщение от", m["from_id"], ":", m["text"][:80]); safe_handle(m["from_id"], m["text"])
        except Exception as e:
            log("переподключение:", e); srv = None; time.sleep(2)


CLI = "--cli" in sys.argv
if __name__ == "__main__":
    if CLI:
        print("Чат с Настеной в консоли. /сброс — начать заново, пустая строка — выход.")
        while True:
            t = input("\nВы: ").strip()
            if not t: break
            handle(1, t)
    else:
        main()
